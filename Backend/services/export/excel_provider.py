import io
from typing import List, Tuple
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from models.expense_model import Expense
from services.export.provider_interface import BaseExportProvider

class ExcelExportProvider(BaseExportProvider):
    def generate(self, expenses: List[Expense]) -> Tuple[bytes, str]:
        wb = Workbook()
        ws = wb.active
        ws.title = "Expenses"
        
        # Header setup
        headers = [
            "Expense Date",
            "Category",
            "Subcategory",
            "Amount",
            "Payment Method",
            "Note",
            "Created At"
        ]
        
        ws.append(headers)
        
        # Style header
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            
        # Write rows
        for exp in expenses:
            date_str = exp.expense_date.isoformat() if exp.expense_date else ""
            created_str = exp.created_at.strftime("%Y-%m-%d %H:%M:%S") if exp.created_at else ""
            
            ws.append([
                date_str,
                exp.category or "",
                exp.subcategory or "",
                exp.amount if exp.amount is not None else 0.0,
                exp.payment_method or "",
                exp.note or "",
                created_str
            ])
            
        # Adjust column widths (basic estimation)
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter # Get the column name
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = adjusted_width
            
        # Save to bytes
        output = io.BytesIO()
        wb.save(output)
        excel_bytes = output.getvalue()
        mime_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        
        return excel_bytes, mime_type
